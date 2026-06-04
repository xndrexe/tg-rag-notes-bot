"""Извлечение текста из документов разных форматов."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ExtractError(Exception):
    """Не удалось извлечь текст из документа."""


def extract_text(path: Path) -> str:
    """Возвращает простой текст документа.

    Поддерживаемые форматы: PDF, DOCX, Markdown (`.md`), txt.
    Для неподдерживаемых расширений бросает `ExtractError`.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".md", ".markdown", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise ExtractError(f"Неподдерживаемое расширение: {suffix}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    if not parts:
        raise ExtractError("PDF не содержит извлекаемого текста (вероятно сканы)")
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    if not parts:
        raise ExtractError("DOCX не содержит текста")
    return "\n\n".join(parts)


def chunk_text(text: str, *, chunk_size: int, chunk_overlap: int) -> list[str]:
    """Делит текст на перекрывающиеся куски заданного размера.

    Разбиение по абзацам — если абзац короче chunk_size, остаётся целым.
    Если длиннее — нарезается по символам с overlap.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size должен быть положительным")
    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap должен быть в [0, chunk_size)")

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    buffer = ""

    for para in paragraphs:
        if len(buffer) + len(para) + 2 <= chunk_size:
            buffer = f"{buffer}\n\n{para}" if buffer else para
            continue
        if buffer:
            chunks.append(buffer)
            buffer = ""
        if len(para) <= chunk_size:
            buffer = para
            continue
        chunks.extend(_split_long(para, chunk_size=chunk_size, overlap=chunk_overlap))
    if buffer:
        chunks.append(buffer)
    return chunks


def _split_long(text: str, *, chunk_size: int, overlap: int) -> list[str]:
    pieces: list[str] = []
    start = 0
    step = chunk_size - overlap
    while start < len(text):
        end = min(start + chunk_size, len(text))
        pieces.append(text[start:end])
        if end == len(text):
            break
        start += step
    return pieces
